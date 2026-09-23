"""Focused renderer regressions; no generated publication outputs required."""

import importlib.util
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt
from PIL import Image
import pytest


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location(
    "publication_document", ROOT / "src" / "publication" / "document.py"
)
document = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(document)


def test_renderer_uses_explicit_paths_and_preserves_format(tmp_path):
    refs = tmp_path / "references.json"
    refs.write_text("[]", encoding="utf-8")
    for name in ("first", "second"):
        source = tmp_path / name
        source.mkdir()
        Image.new("RGB", (10, 10), "white").save(source / "figure.png")
        manuscript = source / "article.md"
        manuscript.write_text(
            f"# {name}\n\nBody text.\n\n![Figure](figure.png)\n",
            encoding="utf-8",
        )
        output = tmp_path / "build" / name / "article.docx"
        assert document.build_docx(manuscript, refs, output) == output
        doc = Document(output)
        assert doc.paragraphs[0].text == name
        assert len(doc.inline_shapes) == 1
        assert doc.styles["Normal"].font.name == "Times New Roman"
        assert doc.styles["Normal"].font.size == Pt(12)
        assert doc.styles["References"].font.size == Pt(9)
        assert doc.styles["References"].paragraph_format.line_spacing == 0.95
        assert abs(doc.sections[0].left_margin - Cm(2)) < 635
    assert Document(tmp_path / "build/first/article.docx").paragraphs[0].text == "first"


def test_contact_sheets_are_scoped_to_output_directory(tmp_path):
    page = tmp_path / "page-01.png"
    Image.new("RGB", (100, 140), "white").save(page)
    output = tmp_path / "contact_sheets"
    output.mkdir()
    stale = output / "contact-sheet-99.png"
    stale.touch()
    sentinel = tmp_path / "contact-sheet-99.png"
    sentinel.touch()
    sheets = document.build_contact_sheets([page], output)
    assert sheets == [output / "contact-sheet-01.png"]
    assert sheets[0].is_file()
    assert not stale.exists()
    assert sentinel.exists()


def test_editorial_comments_hidden_without_removing_surrounding_content(tmp_path):
    text = '# Title\n\nBefore <!-- editorial: transfer-candidate TC-01 --> after.\n\n<!-- editorial:\nmultiline note\n-->\n\n| A | B |\n|---|---|\n| 1 | 2 |\n'
    manuscript = tmp_path / 'article.md'
    manuscript.write_text(text)
    refs = tmp_path / 'refs.json'
    refs.write_text('[]')
    output = tmp_path / 'article.docx'
    document.build_docx(manuscript, refs, output)
    rendered = Document(output)
    assert [p.text for p in rendered.paragraphs if p.text] == ['Title', 'Before after.']
    assert len(rendered.tables) == 1
    assert 'editorial' not in document.publication_text(text)
    assert document.publication_text('Unchanged text.') == 'Unchanged text.'


@pytest.mark.parametrize('text', ['Text <!-- missing end', 'Text -->'])
def test_unclosed_editorial_comment_fails(text):
    with pytest.raises(ValueError):
        document.publication_text(text)
