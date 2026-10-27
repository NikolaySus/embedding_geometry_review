import json
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from publication import conference


def test_conference_typography_and_bibliography(tmp_path):
    source=tmp_path/'source.md'
    source.write_text('# Title\n\nText — text [1].\n\n{{REFERENCES}}',encoding='utf-8')
    out=conference.build_docx(source,tmp_path/'test.docx')
    doc=Document(out)
    assert doc.styles['Normal'].font.size.pt==12
    assert doc.styles['Normal'].paragraph_format.line_spacing==1.5
    assert all(abs(getattr(doc.sections[0],k).cm-2)<.01 for k in ['left_margin','right_margin','top_margin','bottom_margin'])
    refs=[p for p in doc.paragraphs if p.style.name=='References']
    assert len(refs)==10
    assert refs[-1].text.startswith('[10] Yu T.')
    assert not any(p.paragraph_format.page_break_before for p in refs)
    assert doc.styles['Caption'].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert 'Text – text [1].' in [p.text for p in doc.paragraphs]
    with ZipFile(out) as package:
        for name in package.namelist():
            if not name.startswith('word/') or not name.endswith('.xml'):
                continue
            root = etree.fromstring(package.read(name))
            for fonts in root.iter(qn('w:rFonts')):
                assert set(fonts.attrib.values()) == {'Times New Roman'}
                assert not any('Theme' in key for key in fonts.attrib)
            for scheme in root.iter(qn('a:fontScheme')):
                assert all(n.get('typeface') == 'Times New Roman' for n in scheme.iter() if 'typeface' in n.attrib)


def test_six_slides_use_all_five_figures():
    slides=json.loads((conference.ROOT/'assets/presentation/slides.json').read_text(encoding='utf-8'))
    assert len(slides)==6
    assert [s['figure'] for s in slides[1:]]==[1,2,3,4,5]
    assert all(len(s['takeaway'].split())<=15 for s in slides[1:])
