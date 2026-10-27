"""Build conference figure assets, the six-slide deck and a separate speaking script."""
import argparse
import hashlib
import json
import os
from pathlib import Path
import subprocess

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Pt

from gallery.bundles import load_bundle
from gallery.surfaces import outside_vertex_labels, front_edge_percentages
from publication import conference, document, unified_assets as shared

ROOT = shared.ROOT
OUT = ROOT / 'build/presentation'


def prepare():
    OUT.mkdir(parents=True, exist_ok=True)
    figures = OUT/'figures'
    data = OUT/'data'
    figures.mkdir(exist_ok=True)
    data.mkdir(exist_ok=True)
    bundle = load_bundle(shared.SETTINGS['experiment'], shared.SETTINGS['model'])
    evidence = shared.evidence(bundle, data)
    spec = json.loads((ROOT/'assets/presentation/slides.json').read_text(encoding='utf-8'))
    manifest = []
    with plt.rc_context({'font.family':'DejaVu Sans', 'font.size':9}):
        for i, fig in enumerate(shared.shared_figures(bundle, data, surface_azimuth=conference.CONFIG['surface_azimuth']), 1):
            labels = [[t.get_text() for t in ax.texts] for ax in fig.axes]
            if i in (2, 3, 5):
                for t in list(fig.texts):
                    t.remove()
                fig.set_size_inches(8.4, 6)
                for j, ax in enumerate(fig.axes):
                    ax.view_init(conference.CONFIG['surface_elevation'], ax.azim)
                    ax.set_position([.10+(j%2)*.49, .52 if j<2 else .02, .40, .47])
                    outside_vertex_labels(ax)
                    # Retain measured-label typography; enlarging it causes roof overlaps.
                    for t in ax.texts:
                        t.set_fontsize(7.5 if t in list(ax.texts)[3:-5] else 8)
            elif i == 4:
                fig.set_size_inches(9, 6.6)
                for j, ax in enumerate(fig.axes):
                    ax.set_position([.09+(j%2)*.49, .60 if j<2 else .16, .38, .32])
                    ax.tick_params(labelsize=10)
                    ax.title.set_fontsize(12)
                    ax.xaxis.label.set_fontsize(10)
            assert labels == [[t.get_text() for t in ax.texts] for ax in fig.axes]
            percentages=[]
            if i in (2,3,5):
                percentages=[[t.get_text() for t in front_edge_percentages(ax,6)] for ax in fig.axes]
                fig.text(.5,.008,'Доли на ближнем ребре: STS / Cls / Ret, %',ha='center',fontsize=8)
            # Text3D tight-bbox measurement misses left-aligned ruler extents.
            kwargs = {} if i in (2,3,5) else dict(bbox_inches='tight',pad_inches=.05)
            fig.savefig(figures/f'figure-{i}.png',dpi=220,**kwargs)
            fig.savefig(figures/f'figure-{i}.pdf',**kwargs)
            manifest.append({'figure':i,'labels':labels,'edge_percentages':percentages,'file':f'figures/figure-{i}.png',
                             'cameras':[{'elevation':ax.elev,'azimuth':ax.azim} for ax in fig.axes if hasattr(ax,'elev')]})
            plt.close(fig)
    (OUT/'figure_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    (OUT/'deck.json').write_text(json.dumps({'author':conference.CONFIG,'slides':spec},ensure_ascii=False,indent=2),encoding='utf-8')
    text = conference.prepare_text((ROOT/'manuscript/talk.md').read_text(encoding='utf-8'))
    text = shared.render_text(text, evidence)
    (OUT/'talk.md').write_text(text,encoding='utf-8')
    doc = Document()
    document.configure_document(doc)
    doc.styles['Normal'].font.size = Pt(14)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.15
    for line in text.splitlines():
        if not line.strip():
            continue
        if line.startswith('# '):
            doc.add_heading(line[2:],0)
        elif line.startswith('## '):
            p=doc.add_heading(line[3:],1)
            p.paragraph_format.page_break_before=True
        else:
            p=doc.add_paragraph(line)
            if line.startswith('['):
                for r in p.runs:
                    r.italic=True
    doc.core_properties.author=conference.CONFIG['author_ru']
    doc.save(OUT/'speaker_script.docx')
    sources=[ROOT/'manuscript/talk.md', ROOT/'assets/presentation/slides.json', ROOT/'configs/conference.json', ROOT/'rules/of.prezentaciya_russkiy_inzhener_shablon.pptx', ROOT/'configs/publication.json', shared.MODEL/'bundle.json', shared.SOURCE/'scores.csv', Path(__file__), ROOT/'src/build_talk.mjs']
    (OUT/'provenance.json').write_text(json.dumps({str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in sources},indent=2),encoding='utf-8')


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--assets-only',action='store_true')
    args=parser.parse_args()
    prepare()
    if not args.assets_only:
        default=Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/node'
        node=Path(os.environ.get('PRESENTATION_NODE',default/'bin/node.exe'))
        subprocess.run([str(node), str(ROOT/'src/build_talk.mjs')],cwd=ROOT,check=True)
        document.export_pdf(OUT/'cross_task_interference_talk.pptx')
    print(f'Conference outputs: {OUT}')


if __name__=='__main__':
    main()
