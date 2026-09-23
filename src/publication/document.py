from __future__ import annotations

import json
import re
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    add_page_number(section.footer.paragraphs[0])

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        normal._element.rPr.rFonts.set(qn(f"w:{attr}"), "Times New Roman")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.widow_control = True

    for name, size in (("Title", 15), ("Heading 1", 13), ("Heading 2", 12), ("Heading 3", 12)):
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            style._element.rPr.rFonts.set(qn(f"w:{attr}"), "Times New Roman")
        border = style._element.pPr.find(qn("w:pBdr"))
        if border is not None:
            style._element.pPr.remove(border)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(9 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(4)
    styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Title"].paragraph_format.first_line_indent = Cm(0)

    caption = styles["Caption"]
    caption.font.name = "Times New Roman"
    caption.font.size = Pt(10)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)
    caption.paragraph_format.space_after = Pt(6)

    if "References" not in styles:
        ref_style = styles.add_style("References", WD_STYLE_TYPE.PARAGRAPH)
    else:
        ref_style = styles["References"]
    ref_style.font.name = "Times New Roman"
    ref_style.font.size = Pt(10)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        ref_style._element.rPr.rFonts.set(qn(f"w:{attr}"), "Times New Roman")
    ref_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    ref_style.paragraph_format.first_line_indent = Cm(-0.6)
    ref_style.paragraph_format.left_indent = Cm(0.6)
    ref_style.paragraph_format.line_spacing = 1.0
    ref_style.paragraph_format.space_after = Pt(3)


INLINE_RE = re.compile(r"(\*\*.*?\*\*|\*.*?\*|\[\d+(?:[–,-]\d+)*\])")


def add_inline(paragraph, text: str) -> None:
    position = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            paragraph.add_run(token)
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def add_paragraph(doc: Document, text: str, *, style: str | None = None) -> None:
    paragraph = doc.add_paragraph(style=style)
    add_inline(paragraph, text)


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        rows.append([cell.strip() for cell in line.strip().strip("|").split("|")])
    return [rows[0], *rows[2:]]


def add_table(doc: Document, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    for row_idx, row in enumerate(rows):
        for col_idx, value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, value)
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(9)
                if row_idx == 0:
                    run.bold = True
            if row_idx == 0:
                set_cell_shading(cell, "D9E2F3")
    set_repeat_table_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_references(doc: Document, refs: list[dict]) -> None:
    for ref in refs:
        paragraph = doc.add_paragraph(style="References")
        paragraph.add_run(f"{ref['id']}. {ref['text']} ")
        run = paragraph.add_run(ref["url"])
        run.font.color.rgb = RGBColor(5, 99, 193)
        run.underline = True


def publication_text(text: str) -> str:
    """Keep source-only editorial comments out of all rendered content."""
    text = re.sub(r"<!--.*?-->", "", text, flags=re.DOTALL)
    if "<!--" in text or "-->" in text:
        raise ValueError("Unclosed or unmatched Markdown editorial comment")
    return text


def build_docx(manuscript: Path, references: Path, output: Path) -> Path:
    refs = json.loads(references.read_text(encoding="utf-8"))
    lines = publication_text(manuscript.read_text(encoding="utf-8")).splitlines()
    doc = Document()
    configure_document(doc)
    doc.styles["References"].font.size = Pt(9)
    doc.styles["References"].paragraph_format.line_spacing = 0.95

    i = 0
    pending_caption = False
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line == "{{REFERENCES}}":
            add_references(doc, refs)
            i += 1
            continue
        if line.startswith("!["):
            match = re.match(r"!\[(.*?)\]\((.*?)\)", line)
            if not match:
                raise ValueError(f"Invalid image directive: {line}")
            image_path = (manuscript.parent / match.group(2)).resolve()
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(str(image_path), width=Inches(6.2))
            pending_caption = True
            i += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            add_table(doc, parse_table(table_lines))
            continue
        if line.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            add_inline(paragraph, line[2:])
            i += 1
            continue
        if line.startswith("## "):
            paragraph = doc.add_heading(line[3:], level=1)
            paragraph.paragraph_format.first_line_indent = Cm(0)
            i += 1
            continue
        if line.startswith("### "):
            paragraph = doc.add_heading(line[4:], level=2)
            paragraph.paragraph_format.first_line_indent = Cm(0)
            i += 1
            continue
        if line.startswith("*Рисунок"):
            paragraph = doc.add_paragraph(style="Caption")
            add_inline(paragraph, line.strip("*"))
            pending_caption = False
            i += 1
            continue
        if line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:])
            i += 1
            continue

        paragraph_lines = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(("#", "|", "![", "- ")) and lines[i] != "{{REFERENCES}}":
            paragraph_lines.append(lines[i].strip())
            i += 1
        text = " ".join(paragraph_lines).replace("  ", " ")
        paragraph = doc.add_paragraph()
        if text.startswith("**Автор:") or text.startswith("**Организация:") or text.startswith("**E-mail:"):
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif text.startswith("**Geometry of") or text.startswith("**Cross-Task"):
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(8)
        add_inline(paragraph, text)

    core = doc.core_properties
    core.title = "Межзадачные компромиссы при дообучении моделей текстовых эмбеддингов"
    core.subject = "Связь переноса качества, геометрии пространства и смешивания обучающих целей"
    core.keywords = "text embeddings, cross-task transfer, geometry, objective mixing, fine-tuning"
    core.language = "ru-RU"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


def export_pdf(docx_path: Path) -> Path:
    libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
    if not libreoffice:
        raise RuntimeError("LibreOffice is required to export PDF")
    subprocess.run(
        [libreoffice, "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)],
        check=True,
    )
    pdf = docx_path.with_suffix(".pdf")
    if not pdf.is_file():
        raise RuntimeError(f"LibreOffice did not produce {pdf}")
    return pdf


def render_pages(pdf_path: Path, pages_dir: Path) -> list[Path]:
    pages_dir.mkdir(parents=True, exist_ok=True)
    for old in pages_dir.glob("page-*.png"):
        old.unlink()
    subprocess.run(
        ["pdftoppm", "-png", "-r", "130", str(pdf_path), str(pages_dir / "page")],
        check=True,
    )
    pages = sorted(pages_dir.glob("page-*.png"))
    if not pages:
        raise RuntimeError("No rendered pages were produced")
    return pages


def build_contact_sheets(pages: list[Path], contacts_dir: Path, per_sheet: int = 6) -> list[Path]:
    contacts_dir.mkdir(parents=True, exist_ok=True)
    for old in contacts_dir.glob("contact-sheet-*.png"):
        old.unlink()
    outputs = []
    for sheet_idx, start in enumerate(range(0, len(pages), per_sheet), start=1):
        subset = pages[start : start + per_sheet]
        thumbs = []
        for page in subset:
            image = Image.open(page).convert("RGB")
            image.thumbnail((520, 735))
            thumbs.append((page.name, image.copy()))
        canvas = Image.new("RGB", (1120, 3 * 790), "#D6D6D6")
        draw = ImageDraw.Draw(canvas)
        for idx, (name, image) in enumerate(thumbs):
            col, row = idx % 2, idx // 2
            x, y = 25 + col * 550, 35 + row * 790
            canvas.paste(image, (x, y + 25))
            draw.text((x, y), name, fill="#111111")
        output = contacts_dir / f"contact-sheet-{sheet_idx:02d}.png"
        canvas.save(output)
        outputs.append(output)
    return outputs
