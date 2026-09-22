from __future__ import annotations

import re
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[2]
MANUSCRIPT = ROOT / "manuscript" / "full.md"
DATA = ROOT / "build" / "full" / "data"
FIGURES = ROOT / "build" / "full" / "figures"
MIXTURE_DATA = ROOT / "build" / "full" / "mixtures" / "data"
MIXTURE_FIGURES = ROOT / "build" / "full" / "mixtures" / "figures"
BUILD = ROOT / "build" / "full"
RENDERS = ROOT / "build" / "full" / "pages"
DOCX = BUILD / "cross_task_interference_full.docx"
PDF = BUILD / "cross_task_interference_full.pdf"


def main() -> None:
    errors: list[str] = []
    text = MANUSCRIPT.read_text(encoding="utf-8")

    expected_figures = {
        "experiment_design.png",
        "overlap_audit.png",
        "family_transfer_heatmap.png",
        "target_vs_non_target.png",
        "selected_task_deltas.png",
        "geometry_trajectories.png",
        "geometry_quality_correlations.png",
    }
    missing_figures = [name for name in expected_figures if not (FIGURES / name).exists()]
    if missing_figures:
        errors.append(f"Missing figures: {missing_figures}")
    for name in expected_figures:
        if name not in text:
            errors.append(f"Figure is not referenced by manuscript: {name}")

    expected_mixture_figures = {
        "simplex_family_maps.png",
        "family_delta_heatmap.png",
        "pairwise_pareto.png",
        "interaction_residual_heatmap.png",
        "simplex_geometry_maps.png",
    }
    missing_mixture_figures = [name for name in expected_mixture_figures if not (MIXTURE_FIGURES / name).exists()]
    if missing_mixture_figures:
        errors.append(f"Missing mixture figures: {missing_mixture_figures}")
    for name in expected_mixture_figures:
        if name not in text:
            errors.append(f"Mixture figure is not referenced by manuscript: {name}")

    required_data = {
        "family_deltas_by_seed.csv",
        "family_transfer_summary.csv",
        "target_vs_non_target.csv",
        "geometry_trajectories.csv",
        "geometry_and_quality_deltas.csv",
        "selected_task_deltas.csv",
        "exact_text_overlaps.csv",
        "clean_sts_pair_counts.csv",
    }
    missing_data = [name for name in required_data if not (DATA / name).exists()]
    if missing_data:
        errors.append(f"Missing derived data: {missing_data}")

    required_mixture_data = {
        "mixture_summary.csv",
        "pairwise_pareto.csv",
        "interaction_residual_summary.csv",
        "geometry_delta_summary.csv",
    }
    missing_mixture_data = [name for name in required_mixture_data if not (MIXTURE_DATA / name).exists()]
    if missing_mixture_data:
        errors.append(f"Missing mixture data: {missing_mixture_data}")

    if not missing_data:
        family = pd.read_csv(DATA / "family_deltas_by_seed.csv")
        if len(family) != 48:
            errors.append(f"Expected 48 family-delta rows, found {len(family)}")
        if set(family["seed"]) != {42, 43, 44}:
            errors.append("Unexpected seed set in family deltas")
        geometry = pd.read_csv(DATA / "geometry_and_quality_deltas.csv")
        if len(geometry) != 12:
            errors.append(f"Expected 12 geometry/quality observations, found {len(geometry)}")

    if not missing_mixture_data:
        mixture = pd.read_csv(MIXTURE_DATA / "mixture_summary.csv")
        if len(mixture) != 23:
            errors.append(f"Expected 23 mixture/control rows, found {len(mixture)}")
        pareto = pd.read_csv(MIXTURE_DATA / "pairwise_pareto.csv")
        if len(pareto) != 22:
            errors.append(f"Expected 22 simplex points, found {len(pareto)}")

    if not DOCX.exists() or not PDF.exists():
        errors.append("DOCX or PDF build artifact is missing")
    else:
        document = Document(DOCX)
        document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        if "Межзадачные компромиссы при дообучении моделей текстовых эмбеддингов" not in document_text:
            errors.append("Experimental title is missing from DOCX")
        if len(document.tables) < 9:
            errors.append(f"Expected at least 9 tables, found {len(document.tables)}")
        page_count = len(PdfReader(PDF).pages)
        if not 14 <= page_count <= 40:
            errors.append(f"Unexpected PDF page count: {page_count}")
        rendered = sorted(RENDERS.glob("page-*.png"))
        if len(rendered) != page_count:
            errors.append("Rendered page count does not match PDF page count")

    forbidden_claims = [
        "полные спектры ковариации были измерены",
        "траектории downstream-качества",
    ]
    for claim in forbidden_claims:
        if claim in text.lower():
            errors.append(f"Unsupported claim found: {claim}")

    for removed in ("Raw MiniLM", "raw_minilm", "нескольких компактных моделей", "две исходные модели"):
        if removed in text:
            errors.append(f"Removed model claim found: {removed}")
    if not missing_data:
        if set(family.model) != {"all_minilm"} or set(geometry.model) != {"all_minilm"}:
            errors.append("Unexpected model in full data")

    required_abstract_sentence = (
        "смешивание обучающих целей ослабляет, но не устраняет выявленные межзадачные компромиссы"
    )
    if required_abstract_sentence not in text:
        errors.append("Approved objective-mixing conclusion is missing from the abstract")

    if errors:
        print("Validation failed:")
        for error in errors:
            print(f"- {error}")
        sys.exit(1)

    words = len(re.findall(r"\b[\w-]+\b", text, flags=re.UNICODE))
    figure_count = len(expected_figures) + len(expected_mixture_figures)
    print(f"Validation passed: {words} words, {figure_count} figures, {len(PdfReader(PDF).pages)} pages")


if __name__ == "__main__":
    main()

