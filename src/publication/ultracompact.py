"""Single-row figures and the independently sized conference manuscript."""

import hashlib
import json
from pathlib import Path
import re
import shutil
from tempfile import TemporaryDirectory

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from pypdf import PdfReader

from gallery.bundles import load_bundle
from gallery.surfaces import outside_vertex_labels, front_edge_percentages
from . import unified_assets as shared
from . import conference

ROOT = shared.ROOT
OUT = ROOT / "build/ultracompact"
FIGURE_IDS = (shared.COMPACT_IDS[0], "2D-row", "3E-row", "4A-row", "5A-row")


def surface_row(fig, outside_labels=True):
    """Reflow vector artists and lower the camera, retaining data and labels."""
    for text in list(fig.texts):
        text.remove()
    fig.set_size_inches(6.2, 1.85)
    for index, ax in enumerate(fig.axes):
        ax.view_init(conference.CONFIG['surface_elevation'], ax.azim)
        ax.set_position([.036 + index * .245, .03, .235, .94])
        if outside_labels:
            outside_vertex_labels(ax)
        for text in ax.texts:
            text.set_fontsize(text.get_fontsize() * .53)
        for line in ax.lines:
            line.set_linewidth(line.get_linewidth() * .53)
        for collection in ax.collections:
            collection.set_linewidth(np.asarray(collection.get_linewidth()) * .53)
            if hasattr(collection, "get_sizes"):
                collection.set_sizes(collection.get_sizes() * .53**2)
    return fig


def trajectory_row(fig):
    fig.set_size_inches(6.2, 1.9)
    for legend in list(fig.legends):
        legend.remove()
    titles = ["Gate: Δ STS", "Gate: Δ Ret", "Effective rank", "kNN preservation"]
    for index, (ax, title) in enumerate(zip(fig.axes, titles)):
        ax.set_position([.065 + index * .245, .30, .19, .55])
        ax.set_title(title, fontsize=6.5, pad=4)
        ax.set_xlabel("Steps", fontsize=6, labelpad=2)
        ax.tick_params(labelsize=5.5, pad=1)
        ax.yaxis.get_offset_text().set_fontsize(5.5)
        for line in ax.lines:
            line.set_linewidth(.8)
            line.set_markersize(2)
    fig.legend(*fig.axes[0].get_legend_handles_labels(), loc="lower center",
               ncol=6, frameon=False, fontsize=6, columnspacing=.8, handlelength=1.8)
    return fig


def main(output=None):
    output = output or OUT
    data, figures = output / "data", output / "figures"
    data.mkdir(parents=True, exist_ok=True)
    figures.mkdir(parents=True, exist_ok=True)
    bundle = load_bundle(shared.SETTINGS["experiment"], shared.SETTINGS["model"])
    # Copy only recorded evidence; never invoke training, evaluation, or gallery builds.
    sources = list(shared.MODEL.glob("*.csv")) + [shared.MODEL / "bundle.json", shared.SOURCE / "runs.csv", shared.SOURCE / "scores.csv", ROOT / "configs/unified_v1.json"] + shared.AUDIT_FILES
    if set(pd.read_csv(shared.SOURCE / "runs.csv").protocol) != {shared.SETTINGS["protocol_hash"]}:
        raise ValueError("Publication protocol mismatch")
    hashes = {}
    for path in sources:
        shutil.copyfile(path, data / path.name)
        hashes[str(path.relative_to(ROOT))] = hashlib.sha256(path.read_bytes()).hexdigest()
    for path in [Path(__file__), ROOT / "src/publication/unified_assets.py", ROOT / "configs/publication.json", ROOT / "manuscript/ultracompact.md", ROOT / "configs/conference.json", ROOT / "src/publication/conference.py", ROOT / "references/references.json"]:
        hashes[str(path.relative_to(ROOT))] = hashlib.sha256(path.read_bytes()).hexdigest()
    values = shared.evidence(bundle, data)
    (data / "provenance.json").write_text(json.dumps({**shared.SETTINGS, "files": hashes}, indent=2))
    manifest = []
    with plt.rc_context({"font.family": "DejaVu Sans", "font.size": 8, "axes.labelsize": 8}):
        plots = shared.shared_figures(bundle, data, surface_azimuth=conference.CONFIG['surface_azimuth'])
        for i, (name, fig) in enumerate(zip(FIGURE_IDS, plots)):
            before = [[t.get_text() for t in ax.texts] for ax in fig.axes]
            if i in (1, 2, 4):
                surface_row(fig)
            elif i == 3:
                trajectory_row(fig)
            after = [[t.get_text() for t in ax.texts] for ax in fig.axes]
            assert before == after, "A numeric surface label was lost"
            percentages = []
            if i in (1,2,4):
                percentages = [[t.get_text() for t in front_edge_percentages(ax,3)] for ax in fig.axes]
                fig.text(.5,.008,'Доли на ближнем ребре: STS / Cls / Ret, %',ha='center',fontsize=4)
            width, height = fig.get_size_inches()
            fig.savefig(figures / f"{name}.pdf")
            fig.savefig(figures / f"{name}.png", dpi=180 if i == 0 else 420)
            manifest.append({"id": name, "width_inches": float(width), "height_inches": float(height),
                             "axes": [list(ax.get_position().bounds) for ax in fig.axes], "texts": after, "edge_percentages": percentages,
                             "cameras": [{"elevation": ax.elev, "azimuth": ax.azim} for ax in fig.axes if hasattr(ax, 'elev')]})
            plt.close(fig)
    (data / "figure_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2))
    return values


def validate():
    source = (ROOT / "manuscript/ultracompact.md").read_text()
    text = (OUT / "manuscript.md").read_text()
    values = json.loads((OUT / "data/evidence.json").read_text())
    with TemporaryDirectory() as directory:
        fresh = shared.evidence(load_bundle(shared.SETTINGS["experiment"], shared.SETTINGS["model"]), Path(directory))
    assert values == fresh, "Stale numerical evidence"
    expected = shared.render_text(conference.prepare_text(source), values).replace("../build/ultracompact/figures/", "figures/")
    assert text == expected
    assert re.findall(r"\{\{(.*?)\}\}", text) == ["REFERENCES"]
    provenance = json.loads((OUT / "data/provenance.json").read_text())
    for path, digest in provenance["files"].items():
        assert hashlib.sha256((ROOT / path).read_bytes()).hexdigest() == digest, path
    doc = Document(OUT / "cross_task_interference_ultracompact.docx")
    pdf = PdfReader(OUT / "cross_task_interference_ultracompact.pdf")
    assert len(doc.inline_shapes) == 5 and len(doc.tables) == 0
    assert doc.styles["Normal"].font.size.pt == 12
    assert doc.styles["Normal"].paragraph_format.line_spacing == 1.5
    for field in ['author_ru', 'author_en', 'email', 'affiliation_ru', 'affiliation_en']:
        assert conference.CONFIG[field] in text
    assert len(pdf.pages) == 6
    for heading in ['Введение', 'Материалы и методы', 'Результаты', 'Обсуждение', 'Заключение', 'Список источников']:
        assert any(p.text == heading for p in doc.paragraphs)
    assert len([p for p in doc.paragraphs if p.style.name == 'References']) == 10
    body = '\n'.join(p.text for p in doc.paragraphs if p.style.name != 'References')
    first_mentions = list(dict.fromkeys(int(n.strip()) for group in re.findall(r'\[(\d+(?:,\s*\d+)*)\]', body) for n in group.split(',')))
    assert first_mentions == list(range(1,11)), first_mentions
    assert all(np.isclose(shape.width / 914400, 6.2) for shape in doc.inline_shapes)
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        assert np.isclose(getattr(doc.sections[0], attr).cm, 2, atol=.01)
    manifest = json.loads((OUT / "data/figure_manifest.json").read_text())
    assert [r["id"] for r in manifest] == list(FIGURE_IDS)
    for item in manifest[1:]:
        axes = item["axes"]
        assert len(axes) == 4
        assert np.allclose([a[1] for a in axes], axes[0][1])
        assert all(axes[i][0] < axes[i+1][0] for i in range(3))
        figure_pdf = PdfReader(OUT / "figures" / f"{item['id']}.pdf")
        assert not figure_pdf.pages[0].images, "Row figures must remain vector graphics"
        if item["id"] != "4A-row":
            assert all(len(labels) == 30 for labels in item["texts"])
    reference_pages = [i for i, page in enumerate(pdf.pages) if "Список источников" in page.extract_text(extraction_mode='layout')]
    assert len(reference_pages) == 1
    for page in pdf.pages:
        if re.search(r"Рисунок\s+\d+\s+[—–-]", page.extract_text()):
            assert len(page.images), "Detached figure caption"
    result = {"references_start_page": reference_pages[0]+1,
              "total_pages": len(pdf.pages), "figures": 5, "tables": 0, "sources": 10,
              "citation_first_appearance": first_mentions}
    (OUT / "validation.json").write_text(json.dumps(result, indent=2))
    print(json.dumps(result, indent=2))
