"""Conference-only typography and metadata; other article variants are unchanged."""
import json
import re
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree
from . import document

ROOT = Path(__file__).resolve().parents[2]
CONFIG = json.loads((ROOT / 'configs/conference.json').read_text(encoding='utf-8'))


def selected_references():
    refs = json.loads((ROOT / 'references/references.json').read_text(encoding='utf-8'))
    by_id = {r['id']: r for r in refs}
    return [dict(by_id[old], id=i+1, source_id=old) for i, old in enumerate(CONFIG['reference_ids'])]


def prepare_text(text):
    text = text.replace('\u2014', '\u2013')
    for name, value in CONFIG.items():
        if isinstance(value, str):
            text = text.replace('{{'+name+'}}', value)
    ids = {old: i+1 for i, old in enumerate(CONFIG['reference_ids'])}
    def citation(match):
        return '['+', '.join(str(ids[int(n.strip())]) for n in match[1].split(','))+']'
    return re.sub(r'\[(\d+(?:,\s*\d+)*)\]', citation, text)


def build_docx(manuscript, output):
    doc = Document()
    document.configure_document(doc)
    for name in ['Normal', 'Title', 'Heading 1', 'Caption', 'References']:
        s = doc.styles[name]
        s.font.name = 'Times New Roman'
        s.font.size = Pt(12)
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.line_spacing = 1.5
        s.paragraph_format.space_before = Pt(0)
        s.paragraph_format.space_after = Pt(0)
        s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.styles['Title'].font.bold = True
    doc.styles['Caption'].font.italic = False
    doc.styles['Caption'].font.bold = False
    doc.styles['Caption'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles['References'].paragraph_format.left_indent = Cm(0)
    doc.styles['References'].paragraph_format.first_line_indent = Cm(0)
    lines = manuscript.read_text(encoding='utf-8').splitlines()
    for line in lines:
        if not line.strip():
            continue
        if line == '{{REFERENCES}}':
            for r in selected_references():
                p = doc.add_paragraph(style='References')
                p.add_run(f"[{r['id']}] {r['text']}")
                if 'DOI:' not in r['text']:
                    p.add_run(' '+r['url'])
            continue
        if line.startswith('!['):
            m = re.fullmatch(r'!\[(.*?)\]\((.*?)\)', line)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.line_spacing = 1
            p.add_run().add_picture(str((manuscript.parent/m[2]).resolve()), width=Inches(6.2))
            continue
        if line.startswith('## '):
            p = doc.add_paragraph(style='Heading 1')
            p.add_run(line[3:])
        elif line.startswith('# '):
            p = doc.add_paragraph(style='Title')
            p.add_run(line[2:])
        elif line.startswith('*Рисунок'):
            p = doc.add_paragraph(style='Caption')
            p.add_run(line.strip('*'))
        else:
            if line.startswith('**Cross-Task'):
                spacer = doc.add_paragraph()
                spacer.paragraph_format.keep_with_next = True
            p = doc.add_paragraph()
            document.add_inline(p, line)
            if line.startswith(('**Cross-Task', 'Gorkunov', 'BMSTU')):
                p.paragraph_format.keep_with_next = True
            if line.startswith(('УДК', 'Горкунов', 'Gorkunov', 'SPIN', 'МГТУ', 'BMSTU', '**Cross-Task')):
                p.paragraph_format.first_line_indent = Cm(0)
    doc.core_properties.author = CONFIG['author_ru']
    doc.core_properties.title = next(line[2:] for line in lines if line.startswith('# '))
    doc.core_properties.language = 'ru-RU'
    # Theme references override literal fonts in Word, notably on headings.
    # Cover all script slots and all styles, including linked character styles.
    for style in doc.styles:
        if hasattr(style, 'font'):
            style.font.name = 'Times New Roman'
    for part in doc.part.package.parts:
        root = getattr(part, 'element', None)
        generic_xml = root is None and str(part.partname).startswith('/word/') and str(part.partname).endswith('.xml')
        if generic_xml:
            root = etree.fromstring(part.blob)
        if root is not None:
            for fonts in root.iter(qn('w:rFonts')):
                fonts.attrib.clear()
                for slot in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
                    fonts.set(qn('w:'+slot), 'Times New Roman')
            for text_node in root.iter(qn('w:t')):
                if text_node.text:
                    text_node.text = text_node.text.replace('\u2014', '\u2013')
            for scheme in root.iter(qn('a:fontScheme')):
                for node in scheme.iter():
                    if 'typeface' in node.attrib:
                        node.set('typeface', 'Times New Roman')
        if generic_xml:
            part._blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
    doc.save(output)
    return output
