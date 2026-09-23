from pathlib import Path
import re

import numpy as np
import pandas as pd
from docx import Document
from pypdf import PdfReader
import json

from .compact_assets import FIGURE_IDS, evidence, interaction_markdown_rows
from .document import publication_text

ROOT = Path(__file__).resolve().parents[2]


def main():
    source = (ROOT/'manuscript/compact.md').read_text()
    text = publication_text(source)
    figures = re.findall(r'!\[[^\]]*\]\(([^)]+)\)',text)
    assert figures == [f'../build/compact/figures/{identity}.png' for identity in FIGURE_IDS]
    for path in figures:
        assert (ROOT/'manuscript'/path).is_file(), path
    assert not re.search(r'Raw MiniLM|raw_minilm|нескольких компактных моделей',text)
    data=ROOT/'build/compact/data'
    trajectory=pd.read_csv(data/'geometry_trajectories.csv')
    geometry=pd.read_csv(data/'geometry_quality.csv')
    assert set(trajectory.model)==set(geometry.model)=={'all_minilm'}
    assert len(trajectory)==48 and len(geometry)==12
    assert np.allclose(trajectory[trajectory.fraction==0].effective_rank,203.610870,atol=.001)
    q=pd.read_csv(data/'quality_geometry_by_seed.csv')
    assert len(q)==69 and set(q.seed)=={42,43,44}
    pdf=PdfReader(ROOT/'build/compact/cross_task_interference_compact.pdf')
    doc=Document(ROOT/'build/compact/cross_task_interference_compact.docx')
    assert len(doc.inline_shapes)==5
    assert len(doc.tables)==6
    for shape in doc.inline_shapes:
        assert np.isclose(shape.width / 914400, 6.2)
    assert re.findall(r'\*Рисунок (\d+) —', text) == list('12345')
    assert len(set(re.findall(r'transfer-candidate (TC-\d+)', source))) == 6
    assert 'editorial:' not in '\n'.join(p.extract_text() for p in pdf.pages)
    details, _ = evidence()
    for item in details['correlations']:
        assert f"r = {item['r']:.3f}".replace('.', ',') in text
        assert f"n = {item['n']}" in text
    for row in interaction_markdown_rows():
        assert row in text, row
    assert ', '.join(details['pareto_fronts']['ag'][:-1]) + ' и ' + details['pareto_fronts']['ag'][-1] in text
    manifest = json.loads((data/'figure_manifest.json').read_text())
    assert [item['id'] for item in manifest] == list(FIGURE_IDS)
    for item in manifest:
        page = PdfReader(ROOT/'build/compact/figures'/f"{item['id']}.pdf").pages[0]
        assert np.isclose(float(page.mediabox.width), item['width_inches']*72)
        assert np.isclose(float(page.mediabox.height), item['height_inches']*72)
    assert len(list((ROOT/'build/compact/pages').glob('page-*.png')))==len(pdf.pages)
    assert 'Raw MiniLM' not in '\n'.join(p.extract_text() for p in pdf.pages)
    words=len(re.findall(r'\b[\w-]+\b',text))
    print(f'Validated: {words} words, {len(pdf.pages)} pages, 5 dense figures, {len(doc.tables)} tables')


if __name__=='__main__':
    main()
